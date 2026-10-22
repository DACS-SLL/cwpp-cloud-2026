"""Genera el DOCX editable del aporte del integrante 2 desde Markdown versionado."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/informe-integrante-2.md"
OUTPUT = ROOT / "docs/entregables/informe-integrante-2.docx"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    properties.append(shd)


def keep_together(paragraph):
    paragraph.paragraph_format.keep_together = True


def add_rich(paragraph, raw):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
    for segment in pattern.split(raw):
        if not segment:
            continue
        if segment.startswith("**") and segment.endswith("**"):
            run = paragraph.add_run(segment[2:-2])
            run.bold = True
        elif segment.startswith("`") and segment.endswith("`"):
            run = paragraph.add_run(segment[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(20, 87, 95)
        elif segment.startswith("[") and "](" in segment:
            label, path = segment[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(7, 111, 107)
            run.underline = True
            if path.startswith("evidencias/") and path.endswith(".json"):
                paragraph.add_run(f" ({path})")
        else:
            paragraph.add_run(segment.replace("  ", " "))


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("CWPP UTS  |  Aporte integrante 2  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor(29, 48, 61)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    title_style = doc.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(21)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(13)
    title_ppr = title_style._element.find(qn("w:pPr"))
    if title_ppr is not None:
        title_border = title_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_ppr.remove(title_border)
    for style_name, size, before in [("Heading 1", 15, 18), ("Heading 2", 12, 13)]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(10, 74, 80)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    add_footer(section)

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_rich(paragraph, line[2:])
            keep_together(paragraph)
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_rich(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_rich(paragraph, line[4:])
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
            if match:
                image_path = SOURCE.parent / match.group(2)
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.add_run().add_picture(str(image_path), width=Inches(5.95))
                keep_together(paragraph)
        elif line.startswith("| "):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            data = [[cell.strip() for cell in row.strip("|").split("|")] for row in rows]
            data = [row for row in data if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in row)]
            if not data:
                continue
            table = doc.add_table(rows=1, cols=len(data[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            table.autofit = True
            for row_index, values in enumerate(data):
                cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    add_rich(paragraph, value)
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
                    if row_index == 0:
                        shade(cell, "173B4B")
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    elif row_index % 2 == 0:
                        shade(cell, "F1F6F5")
            header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            header_tr_pr.append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_rich(paragraph, line[2:])
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.right_indent = Cm(0.4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(11)
            add_rich(paragraph, line[2:])
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(50, 82, 93)
                run.italic = True
        else:
            paragraph = doc.add_paragraph()
            add_rich(paragraph, line)
            if line.startswith("**Figura"):
                paragraph.paragraph_format.space_after = Pt(11)
                paragraph.paragraph_format.keep_with_next = False

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
